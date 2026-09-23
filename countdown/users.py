"""The user directory.

Who a person is - name, personal number, department - is held in a second
workbook, linked from countdown.txt. The app never asks: it takes the personal
number from the Windows logon name (R3) and looks the row up.

Writing back is only possible when the link is a path. An anonymous SharePoint
link can be read and nothing more, so the maintenance window says so rather than
offering an edit that cannot be saved.
"""
import io
import logging
import os
import re

from . import paths, table

log = logging.getLogger(__name__)

COLUMNS = [
    ("name", ["name", "full name", "user", "user name"]),
    ("personal_number", ["personal number", "personal no", "personal_number",
                         "number", "id", "employee number"]),
    ("department", ["department", "dept"]),
]
REQUIRED = ("personal_number", "department")
HEADERS = ["Name", "Personal number", "Department"]


class UsersError(Exception):
    """The directory could not be read, or holds none of the expected columns."""


class User:
    __slots__ = ("name", "personal_number", "department", "row_number")

    def __init__(self, name, personal_number, department, row_number=0):
        self.name = str(name or "").strip()
        self.personal_number = normalise_number(personal_number)
        self.department = str(department or "").strip()
        self.row_number = row_number

    def __repr__(self):
        return f"<User {self.personal_number} {self.name} {self.department}>"

    def __eq__(self, other):
        return (isinstance(other, User)
                and self.personal_number == other.personal_number
                and self.name == other.name
                and self.department == other.department)


def normalise_number(value) -> str:
    """Personal numbers are compared as digits, however the cell was typed.

    Excel turns a number into a float, so 8123456 can arrive as "8123456.0",
    and people pad with spaces or write `iaf\\8123456`.
    """
    text = str(value if value is not None else "").strip()
    if not text:
        return ""
    text = text.split("\\")[-1].split("/")[-1]
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".")[0]
    digits = re.sub(r"\D", "", text)
    return digits or text


def _map_headers(cells):
    seen = {table._normalise(cell): index for index, cell in enumerate(cells)
            if str(cell or "").strip()}
    mapping = {}
    for field, spellings in COLUMNS:
        for spelling in spellings:
            if spelling in seen:
                mapping[field] = seen[spelling]
                break
    missing = [f for f in REQUIRED if f not in mapping]
    if missing:
        raise UsersError(
            "the user list is missing required column(s): " + ", ".join(missing)
            + "; headers found: " + ", ".join(sorted(seen)))
    return mapping


def parse(payload: bytes):
    """Return (users, skipped). A row with no personal number is skipped."""
    reader = table._rows_from_xlsx if payload[:2] == b"PK" else table._rows_from_csv
    raw_rows = list(reader(payload))

    header_index, mapping, first_error = None, None, None
    for index, cells in enumerate(raw_rows):
        if not any(str(c or "").strip() for c in cells):
            continue
        try:
            mapping = _map_headers(cells)
        except UsersError as exc:
            if first_error is None:
                first_error = exc
            if index < 10:
                continue
            raise
        header_index = index
        break
    if mapping is None:
        raise first_error or UsersError("the user list holds no header row")

    users, skipped = [], []
    for offset, cells in enumerate(raw_rows[header_index + 1:], start=header_index + 2):
        if not any(str(c or "").strip() for c in cells):
            continue
        values = {field: (cells[col] if col < len(cells) else "")
                  for field, col in mapping.items()}
        user = User(values.get("name"), values.get("personal_number"),
                    values.get("department"), offset)
        if not user.personal_number or not user.department:
            skipped.append({"row": offset,
                            "name": str(values.get("name") or "").strip(),
                            "personal_number": user.personal_number,
                            "department": user.department})
            continue
        users.append(user)
    return users, skipped


def load(source: str):
    if not source:
        raise UsersError("no user list is configured in countdown.txt")
    try:
        payload = table.download(source)
    except table.TableError as exc:
        raise UsersError(str(exc)) from exc
    return parse(payload)


def find(users, personal_number: str):
    """The user with this personal number, or None."""
    wanted = normalise_number(personal_number)
    if not wanted:
        return None
    for user in users:
        if user.personal_number == wanted:
            return user
    return None


# --- writing back ---------------------------------------------------------

def writable_path(source: str):
    """The path the directory can be saved to, or None when it is read only."""
    path = table.resolve_local(source)
    if path is None:
        return None
    folder = os.path.dirname(path) or paths.app_dir()
    if os.path.exists(path):
        return path if os.access(path, os.W_OK) else None
    return path if os.access(folder, os.W_OK) else None


def save(source: str, users) -> None:
    """Write the directory back, keeping the sheet's own look where there is one."""
    import openpyxl

    path = writable_path(source)
    if path is None:
        raise UsersError(
            "this user list cannot be edited from here. A SharePoint link can "
            "only be read; point users_url at a file path to edit it.")

    ordered = sorted(users, key=lambda u: (u.department.casefold(), u.name.casefold()))

    if os.path.exists(path) and path.lower().endswith((".xlsx", ".xlsm")):
        book = openpyxl.load_workbook(path)
        sheet = book[book.sheetnames[0]]
        header_index, mapping = _locate_header(sheet)
        # Clear the old rows, then write the new ones into the same columns.
        if sheet.max_row > header_index:
            sheet.delete_rows(header_index + 1, sheet.max_row - header_index)
        for offset, user in enumerate(ordered, start=header_index + 1):
            for field, column in mapping.items():
                sheet.cell(row=offset, column=column + 1,
                           value=getattr(user, field))
    else:
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Users"
        sheet.append(HEADERS)
        for user in ordered:
            sheet.append([user.name, user.personal_number, user.department])
        for column, width in zip("ABC", (26, 18, 20)):
            sheet.column_dimensions[column].width = width

    directory = os.path.dirname(path)
    if directory:
        os.makedirs(directory, exist_ok=True)
    book.save(path)
    log.info("wrote %d users to %s", len(ordered), path)


def _locate_header(sheet):
    for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if not any(str(c or "").strip() for c in row):
            continue
        try:
            return index, _map_headers(list(row))
        except UsersError:
            if index < 10:
                continue
            raise
    raise UsersError("the user list holds no header row")


def add(users, name, personal_number, department):
    """A copy of `users` with this person added. Raises on a bad entry."""
    number = normalise_number(personal_number)
    if not str(name or "").strip():
        raise UsersError("Enter a name.")
    if not number:
        raise UsersError("Enter a personal number.")
    if not str(department or "").strip():
        raise UsersError("Choose a department.")
    if find(users, number):
        raise UsersError(f"Personal number {number} is already on the list.")
    return list(users) + [User(name, number, department)]


def remove(users, personal_number):
    number = normalise_number(personal_number)
    return [u for u in users if u.personal_number != number]
