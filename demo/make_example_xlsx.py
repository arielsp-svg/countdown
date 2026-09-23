"""Write demo/example-ro-table.xlsx, a sample of the table the app reads.

Dates are generated relative to the day it is run, so the example is always
live: re-run it whenever the committed copy has aged.

    python demo/make_example_xlsx.py

Every row is here to show one behaviour from the requirements, and the file is
laid out the way a maintained table would be: a title row above the header,
real date cells rather than text, and the columns in no particular order, since
the app matches them by name.
"""
import os
import sys
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from countdown.engine import add_months  # noqa: E402

TODAY = date.today()

HEADERS = ["Department", "Type of system", "Platform", "Nav system",
           "Receiver model", "Using TOD", "Comments", "RO date"]

# (department, type, platform, nav, receiver, tod, comment, RO date)
ROWS = [
    ("Avionics", "Radar", "F-16 Barak", "INS-4", "R-200B", "yes",
     "Inside 6 months: alerts weekly", add_months(TODAY, 4)),
    ("Avionics", "Comms", "F-15 Ra'am", "INS-2", "R-310", "no",
     "Inside 12 months: alerts monthly", add_months(TODAY, 10)),
    ("Avionics", "Nav", "C-130 Shimshon", "GPS-7", "R-105", "yes",
     "Beyond 12 months: no alert yet", add_months(TODAY, 18)),
    ("Avionics", "Radar", "Apache AH-64", "INS-1", "R-55", "no",
     "RO date already passed: no alert", add_months(TODAY, -3)),
    ("Avionics", "Comms", "Blackhawk", "GPS-3", "R-70", "yes",
     "RO date cell holds text: row is skipped", "to be confirmed"),
    ("Avionics", "Nav", "Beechcraft 200", "GPS-1", "R-18", "no",
     "RO date cell empty: row is skipped", None),
    ("Logistics", "Radar", "Hercules loader", "INS-9", "R-12", "no",
     "Another department: filtered out", add_months(TODAY, 2)),
    ("Logistics", "Comms", "Ground station B", "GPS-5", "R-88", "yes",
     "Another department: filtered out", add_months(TODAY, 7)),
    ("Maintenance", "Nav", "Test bench 4", "INS-6", "R-41", "no",
     "Another department: filtered out", add_months(TODAY, 5)),
]


def main(path):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "RO dates"

    # A title row above the header: the app skips down to the real header.
    sheet.append(["Systems RO tracking", None, None, None, None, None, None, None])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append([])
    sheet.append(HEADERS)

    header_row = 3
    fill = PatternFill("solid", fgColor="E8EEF7")
    for column in range(1, len(HEADERS) + 1):
        cell = sheet.cell(row=header_row, column=column)
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")

    for row in ROWS:
        sheet.append(list(row))

    # The RO date column holds real dates, so how Excel displays them does not
    # change what the app reads.
    for index in range(len(ROWS)):
        cell = sheet.cell(row=header_row + 1 + index, column=len(HEADERS))
        if isinstance(ROWS[index][-1], date):
            cell.number_format = "DD/MM/YYYY"

    widths = [14, 16, 18, 13, 16, 11, 40, 13]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)

    book.save(path)
    print(f"wrote {path}")


if __name__ == "__main__":
    default = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                           "example-ro-table.xlsx")
    main(sys.argv[1] if len(sys.argv) > 1 else default)
